"""Gemeinsamer, lokaler Datei-Upload. Erst Vorschau, dann ausdrückliches Speichern.

Keine Trainingsberechnung, keine automatische OCR und keine externen Dienste.
Die jeweilige App behält ihre eigene Datenprüfung und Speichertransaktion.
"""
import csv
import hashlib
import io
import json
import re
import unicodedata
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

FIELD_MAX_ROWS = 1000
UPLOAD_LABEL = 'Kaderdaten hochladen (Sicherung, Excel, PDF oder CSV)'
UPLOAD_HELP = 'Datei auswählen, erkannte Personen prüfen und übernehmen. JSON stellt die Sicherung wieder her; Tabellen liefern die darin enthaltenen Angaben.'


def field_text(value):
    return '' if value is None else str(value).strip()


def header_key(value):
    return re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKD', field_text(value)).casefold())


def upload_kind(raw, filename):
    if not raw:
        raise ValueError('Die Datei ist leer. Bitte eine gespeicherte Sicherung oder Tabelle auswählen.')
    if len(raw) > 200_000_000:
        raise ValueError('Die Datei ist größer als 200 MB.')
    suffix = Path(filename).suffix.lower()
    start = raw.lstrip(b'\xef\xbb\xbf \t\r\n')
    if suffix == '.json' or start.startswith((b'{', b'[')):
        return 'backup'
    if suffix == '.pdf' or start.startswith(b'%PDF-'):
        return 'pdf'
    if suffix in ('.xlsx', '.xlsm', '.ods', '.csv', '.tsv', '.txt'):
        return 'table'
    raise ValueError('Bitte eine JSON-Sicherung, Excel-Datei (.xlsx), ODS, CSV oder PDF-Tabelle auswählen. Alte .xls-Dateien zuerst als .xlsx speichern.')


def strict_json(raw):
    def pairs(items):
        result = {}
        for key, value in items:
            if key in result:
                raise ValueError('Die Sicherung enthält einen doppelten Eintrag: '+key)
            result[key] = value
        return result
    def invalid(value):
        raise ValueError('Ungültiger Zahlenwert in der Sicherung: '+value)
    return json.loads(raw, object_pairs_hook=pairs, parse_constant=invalid)


def read_pdf_table(raw):
    """Read ruled, text-based tables; do not guess from a scan or prose.

    Every nonblank page must provide a compatible table. This avoids silently
    importing only the readable part of a mixed scanned/digital document.
    """
    if len(raw) > 10_000_000:
        raise ValueError('Bitte höchstens 10 MB je PDF hochladen.')
    try:
        import pdfplumber
    except ImportError:
        raise ValueError('Für PDF-Tabellen bitte die mitgelieferte requirements.txt aktualisieren.') from None
    result, canonical = [], None
    try:
        with pdfplumber.open(io.BytesIO(raw)) as document:
            if not 1 <= len(document.pages) <= 30:
                raise ValueError('Bitte eine PDF mit 1 bis 30 Seiten verwenden.')
            for page_number, page in enumerate(document.pages, 1):
                if not page.chars and not page.images and not page.lines and not page.rects and not page.curves:
                    continue
                if not page.chars:
                    raise ValueError(f'Seite {page_number} ist ein Scan oder Foto. Eine automatische Texterkennung ist noch nicht angeschlossen. Bitte die ursprüngliche Excel-/CSV-Datei oder die JSON-Sicherung verwenden.')
                tables = page.extract_tables()
                accepted = False
                for table in tables:
                    table = [[field_text(cell).replace('\n', ' ') for cell in row] for row in table]
                    index = next((i for i,row in enumerate(table[:10]) if any(header_key(v) in ('name','vorname','nachname') for v in row)), None)
                    if index is None:
                        # A table without a repeated header cannot be assigned safely.
                        if any(any(cell for cell in row) for row in table):
                            raise ValueError(f'Seite {page_number}: Eine Tabelle hat keine eindeutige Namensüberschrift. Bitte die Kopfzeile auf jeder PDF-Seite mit ausgeben.')
                        continue
                    headers = [header_key(v) for v in table[index]]
                    if canonical is None:
                        canonical = headers
                        result = table[:index+1]
                    elif headers != canonical:
                        raise ValueError('Die PDF enthält unterschiedliche Tabellen. Bitte je Kader-/Testtabelle eine eigene Datei verwenden.')
                    for row in table[index+1:]:
                        if [header_key(v) for v in row] == canonical:
                            continue
                        if any(row):
                            result.append(row)
                    accepted = True
                if not accepted:
                    raise ValueError(f'Seite {page_number}: Keine eindeutig lesbare Tabelle mit Zelllinien und Namensspalte gefunden. Bitte die ursprüngliche Excel-/CSV-Datei oder JSON-Sicherung verwenden.')
                if len(result) > FIELD_MAX_ROWS+20 or any(len(row)>40 for row in result):
                    raise ValueError('Bitte höchstens 1000 Personen und 40 Spalten verwenden.')
    except ValueError:
        raise
    except Exception:
        raise ValueError('Diese PDF konnte nicht gelesen werden. Bitte eine unverschlüsselte PDF-Tabelle oder die Originaltabelle verwenden.') from None
    if canonical is None or len(result)<2:
        raise ValueError('Die PDF enthält keine ausgefüllte Kader-/Testtabelle.')
    return result


def backup_people(data):
    if isinstance(data.get('athletes'), dict):
        return {str(key):record for key,record in data['athletes'].items()}
    return {sport+' / '+name:dict(record, name=name)
            for sport,people in data.items() if isinstance(people,dict)
            for name,record in people.items() if isinstance(record,dict)}


def render_backup_preview(raw, current, revision, decode, merge, save, done, key='kader_backup'):
    """Callbacks validate, construct candidate, save with revision, and refresh UI.

    save(candidate, revision) must reject a stale persistent revision and return
    the new revision; done(candidate, new_revision) owns the app session update.
    """
    import streamlit as st
    import pandas as pd
    try:
        strict_json(raw)
        decoded = decode(raw)
        candidate = merge(decoded, current)
        people = backup_people(candidate)
        old = backup_people(current)
    except (ValueError, TypeError, KeyError, UnicodeError) as exc:
        st.error('Sicherung nicht übernommen: '+str(exc))
        return
    st.caption('JSON-Sicherung erkannt. Enthaltene Profile, Messwerte und Trainingsverläufe werden gemeinsam wiederhergestellt.')
    st.write(f'**Nach der Übernahme: {len(people)} gespeicherte Personen.**')
    if people:
        st.dataframe(pd.DataFrame([{'Name':p.get('name',identity),
            'Gespeicherte Einheiten':len(p.get('einheitenprotokoll',{})),
            'Archivierte Makrozyklen':len(p.get('makrozyklen',{}))}
            for identity,p in people.items()]),hide_index=True,width='stretch')
    removed = [old[k].get('name',k) for k in old.keys()-people.keys()]
    if removed:
        st.warning('Diese bisherigen Personen sind danach nicht mehr im Kader: '+', '.join(sorted(removed)))
    unchanged = candidate == current
    if unchanged:
        st.success('Diese Sicherung ist bereits im Kader vorhanden. Keine erneute Übernahme erforderlich.')
        return
    signature = hashlib.sha256(raw).hexdigest()[:16]+'_'+str(revision)
    confirm = st.checkbox('Vorhandenen Kader durch diese Sicherung ersetzen',key=key+'_confirm_'+signature) if old or current.get('plans') or current.get('legacy_archive') else True
    if old:
        st.caption('Die Vorschau zeigt den Stand nach der Wiederherstellung. Vorhandene Angaben in den enthaltenen Kaderbereichen werden ersetzt.')
    if st.button('Kader übernehmen',type='primary',disabled=not confirm,key=key+'_save_'+signature):
        try:
            # Revalidate the exact uploaded bytes at the point of commitment.
            checked = merge(decode(raw),current)
            new_revision = save(checked,revision)
        except Exception as exc:
            st.error('Nicht übernommen: '+str(exc))
            return
        done(checked,new_revision)


def read_table_cells(data, filename):
    if len(data) > 10_000_000:
        raise ValueError('Bitte höchstens 10 MB je Tabelle hochladen.')
    suffix = Path(filename).suffix.lower()
    limit = FIELD_MAX_ROWS + 20
    if suffix == '.pdf' or data.lstrip().startswith(b'%PDF-'):
        return read_pdf_table(data)
    if suffix in ('.csv', '.tsv', '.txt'):
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('cp1252')
        if text.lower().startswith('sep='):
            text = text.split('\n', 1)[1]
        first = text.splitlines()[0] if text.splitlines() else ''
        sep = ';' if ';' in first else '\t' if '\t' in first else ','
        table = list(csv.reader(io.StringIO(text), delimiter=sep))
    elif suffix in ('.xlsx', '.xlsm', '.ods'):
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 50_000_000:
                raise ValueError('Die entpackte Tabelle ist zu groß.')
            if suffix == '.ods':
                xml = archive.read('content.xml')
                if b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper():
                    raise ValueError('Diese XML-Struktur wird nicht unterstützt.')
                ns = {'t':'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
                      'o':'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
                      'x':'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}
                root = ET.fromstring(xml)
                sheets = root.findall('o:body/o:spreadsheet/t:table', ns)
                if len(sheets) != 1:
                    raise ValueError('ODS: bitte genau ein Tabellenblatt verwenden.')
                def q(prefix, key):
                    return '{'+ns[prefix]+'}'+key
                def physical_rows(parent):
                    for child in parent:
                        if child.tag == q('t', 'table-row'):
                            yield child
                        elif child.tag in {q('t', t) for t in ('table-header-rows','table-rows','table-row-group')}:
                            yield from physical_rows(child)
                table = []
                row_position = 0
                for row in physical_rows(sheets[0]):
                    repeat = int(row.get(q('t','number-rows-repeated'), '1'))
                    if repeat < 1:
                        raise ValueError('Ungültige Zeilenwiederholung.')
                    values = []
                    col = 0
                    for cell in row:
                        if cell.tag not in {q('t','table-cell'),q('t','covered-table-cell')}:
                            continue
                        count = int(cell.get(q('t','number-columns-repeated'), '1'))
                        if count < 1:
                            raise ValueError('Ungültige Spaltenwiederholung.')
                        kind = cell.get(q('o','value-type'))
                        value = cell.get(q('o','date-value')) if kind == 'date' else cell.get(q('o','value')) if kind in ('float','percentage','currency') else '\n'.join(''.join(p.itertext()) for p in cell.findall('x:p', ns))
                        formula = cell.get(q('t','formula'))
                        if formula:
                            value = {'formula':formula}
                        if value and col + count > 40:
                            raise ValueError('Bitte höchstens 40 Spalten verwenden.')
                        values += [value] * min(count, max(0, 40-col))
                        col += count
                    if any(v not in ('', None) for v in values):
                        if row_position + repeat > limit:
                            raise ValueError('Zu viele Tabellenzeilen.')
                        table += [[None]*40 for _ in range(row_position-len(table))]
                        table += [list(values) for _ in range(repeat)]
                    row_position += repeat
            else:
                from openpyxl import load_workbook
                book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
                try:
                    sheets = [s for s in book.worksheets if s.sheet_state == 'visible']
                    if len(sheets) != 1:
                        raise ValueError('Excel: bitte genau ein sichtbares Tabellenblatt verwenden.')
                    sheet = sheets[0]
                    if (sheet.max_column or 0) > 40 or (sheet.max_row or 0) > limit:
                        raise ValueError('Zu viele Zeilen oder Spalten in der Excel-Datei.')
                    table = [[{'formula':c.value} if c.data_type == 'f' else c.value for c in row]
                             for row in sheet.iter_rows(max_row=sheet.max_row or limit, max_col=sheet.max_column or 40)]
                finally:
                    book.close()
    else:
        raise ValueError('Bitte Excel (.xlsx), ODS, CSV oder eine lesbare PDF-Tabelle hochladen.')
    if len(table) > limit or any(len(row) > 40 for row in table):
        raise ValueError('Bitte höchstens 1000 Personen und 40 Spalten verwenden.')
    return table
